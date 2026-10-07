"""Parse downloaded source files into clean {kommun code: value} mappings.

Each parser checks that the file looks the way we expect and raises
SourceError with a readable message when it does not, instead of guessing.
"""

import csv
import io
import json
import re
from pathlib import Path

KOMMUN_CODE = re.compile(r"^\d{4}$")
SCB_REGION = re.compile(r"^\s*(\d{4})\s+\S")


class SourceError(ValueError):
    pass


# --- Öppna jämförelser: Trygghet och säkerhet (crime per 1,000 inhabitants) ---

CRIME_SHEET = "Indikatorer i bostavsordning"  # sic, as named in the workbook
CRIME_COLUMNS = {
    # indicator slug -> text that must appear in that column's header
    "violent-crime": "anmälda våldsbrott per 1 000 invånare",
    "theft": "anmälda stöld- och tillgreppsbrott per 1 000 invånare",
    "vandalism": "anmälda brott om skadegörelse per 1 000 invånare",
}


def _norm_header(text):
    # The workbook mixes normal and non-breaking spaces in "1 000".
    return " ".join(str(text or "").replace("\xa0", " ").split()).lower()


def parse_crime_workbook(path):
    """Return {slug: {code: value}} for the three crime indicators."""
    import openpyxl  # only needed when converting, not at runtime

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if CRIME_SHEET not in workbook.sheetnames:
        raise SourceError(f"sheet {CRIME_SHEET!r} not found; sheets are {workbook.sheetnames}")
    rows = list(workbook[CRIME_SHEET].iter_rows(values_only=True))

    header_index = next((i for i, row in enumerate(rows) if row and row[0] == "K-kod"), None)
    if header_index is None:
        raise SourceError("no header row starting with 'K-kod' found")
    header = [_norm_header(cell) for cell in rows[header_index]]

    columns = {}
    for slug, needle in CRIME_COLUMNS.items():
        matches = [i for i, text in enumerate(header) if needle in text]
        if len(matches) != 1:
            raise SourceError(f"expected one column containing {needle!r}, found {len(matches)}")
        columns[slug] = matches[0]

    result = {slug: {} for slug in CRIME_COLUMNS}
    for row in rows[header_index + 1 :]:
        code = str(row[0]).strip() if row and row[0] is not None else ""
        if not KOMMUN_CODE.match(code):
            continue  # summary rows, blank rows, county headings
        for slug, col in columns.items():
            value = row[col] if col < len(row) else None
            if isinstance(value, (int, float)):
                result[slug][code] = float(value)
    for slug, values in result.items():
        if not values:
            raise SourceError(f"no numeric values found for {slug}")
    return result


# --- SCB Statistikdatabasen: house prices per kommun (CSV export) ---


def _decode(raw):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SourceError("file is neither UTF-8 nor Windows-1252 text")


def _parse_number(text):
    cleaned = str(text).strip().strip('"').replace("\xa0", "").replace(" ", "").replace(",", ".")
    if cleaned in {"", "..", ".", "-"}:
        return None  # SCB uses ".." for missing / confidential
    try:
        return float(cleaned)
    except ValueError:
        return None


def parse_scb_region_csv(path):
    """Return {code: value} from an SCB table exported as CSV.

    Expects one row per region with the region cell formatted like
    "0114 Upplands Väsby" and a single value column (the last column).
    Rows for län or the whole country (two-digit or "00" codes) are skipped.
    """
    text = _decode(Path(path).read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    values = {}
    for row in csv.reader(io.StringIO(text), dialect):
        if not row:
            continue
        match = SCB_REGION.match(row[0])
        if not match:
            continue
        number = _parse_number(row[-1])
        if number is not None:
            if match.group(1) in values:
                raise SourceError(f"kommun {match.group(1)} appears twice; export one year and one measure only")
            values[match.group(1)] = number
    if not values:
        raise SourceError('no rows like "0114 Upplands Väsby,<value>" found; check the export format')
    return values


# --- Committed data files (data/) ---


def write_values_csv(path, values):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(["code", "value"])
        for code in sorted(values):
            writer.writerow([code, f"{values[code]:g}"])


def read_values_csv(path):
    values = {}
    with Path(path).open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != ["code", "value"]:
            raise SourceError(f"{path}: header must be 'code,value', got {reader.fieldnames}")
        for line, row in enumerate(reader, start=2):
            code = row["code"]
            if not KOMMUN_CODE.match(code or ""):
                raise SourceError(f"{path}:{line}: invalid kommun code {code!r}")
            if code in values:
                raise SourceError(f"{path}:{line}: duplicate kommun code {code}")
            try:
                values[code] = float(row["value"])
            except (TypeError, ValueError):
                raise SourceError(f"{path}:{line}: value {row['value']!r} is not a number") from None
    return values


def read_kommuner_csv(path):
    kommuner = {}
    with Path(path).open(newline="", encoding="utf-8") as f:
        for line, row in enumerate(csv.DictReader(f), start=2):
            if not KOMMUN_CODE.match(row.get("code") or "") or not row.get("name"):
                raise SourceError(f"{path}:{line}: need a four-digit code and a name")
            kommuner[row["code"]] = row["name"]
    return kommuner


SLUG = re.compile(r"^[a-z0-9-]+$")
DIMENSION_FIELDS = {"slug", "name", "description"}
INDICATOR_FIELDS = {"slug", "name", "unit", "year", "lower_is_better", "source", "source_url", "dimension", "range"}
# Every indicator has exactly one of: a values file, or a "derived" formula.
# A values file may say where it is imported from ("kolada").
OPTIONAL_INDICATOR_FIELDS = {"file", "kolada", "derived"}
KOLADA_KPI = re.compile(r"^[NU]\d{5}$")


def read_dimensions_json(path):
    dimensions = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(dimensions, list) or not dimensions:
        raise SourceError(f"{path}: expected a non-empty JSON list")
    seen = set()
    for i, item in enumerate(dimensions):
        if not isinstance(item, dict) or set(item) != DIMENSION_FIELDS:
            raise SourceError(f"{path}[{i}]: fields must be exactly {sorted(DIMENSION_FIELDS)}")
        if not SLUG.match(item["slug"]) or item["slug"] in seen:
            raise SourceError(f"{path}[{i}]: slug must be unique lowercase-with-dashes")
        seen.add(item["slug"])
    return dimensions


def _check_range(value):
    return (
        isinstance(value, list)
        and len(value) == 2
        and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in value)
        and value[0] < value[1]
    )


def read_indicators_json(path, dimension_slugs=None):
    """Validate data/indicators.json. Pass dimension_slugs to check each indicator's dimension exists."""
    indicators = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(indicators, list):
        raise SourceError(f"{path}: expected a JSON list")
    slugs = [item.get("slug") for item in indicators if isinstance(item, dict)]
    seen = set()
    for i, item in enumerate(indicators):
        where = f"{path}[{i}]"
        if (
            not isinstance(item, dict)
            or not INDICATOR_FIELDS <= set(item) <= INDICATOR_FIELDS | OPTIONAL_INDICATOR_FIELDS
        ):
            raise SourceError(f"{where}: needs {sorted(INDICATOR_FIELDS)}, may add {sorted(OPTIONAL_INDICATOR_FIELDS)}")
        if not SLUG.match(item["slug"]) or item["slug"] in seen:
            raise SourceError(f"{where}: slug must be unique lowercase-with-dashes")
        if not isinstance(item["lower_is_better"], bool):
            raise SourceError(f"{where}: lower_is_better must be true/false")
        if item["year"] is not None and not isinstance(item["year"], int):
            raise SourceError(f"{where}: year must be an integer or null (not imported yet)")
        if not _check_range(item["range"]):
            raise SourceError(f"{where}: range must be [min, max] with min < max")
        if dimension_slugs is not None and item["dimension"] not in dimension_slugs:
            raise SourceError(f"{where}: unknown dimension {item['dimension']!r}")
        if ("file" in item) == ("derived" in item):
            raise SourceError(f"{where}: needs either 'file' or 'derived' (computed when loading), not both")
        if "kolada" in item and "derived" in item:
            raise SourceError(f"{where}: an indicator is either imported from Kolada or derived, not both")
        if "kolada" in item:
            kolada = item["kolada"]
            if not isinstance(kolada, dict) or not kolada or not set(kolada) <= {"kpi", "search"}:
                raise SourceError(f"{where}: kolada needs 'kpi' and/or 'search'")
            if "kpi" in kolada and not KOLADA_KPI.match(str(kolada["kpi"])):
                raise SourceError(f"{where}: kolada.kpi {kolada['kpi']!r} is not a Kolada id like N00900")
            if "search" in kolada and not (isinstance(kolada["search"], str) and kolada["search"].strip()):
                raise SourceError(f"{where}: kolada.search must be a non-empty string")
        if "derived" in item:
            derived = item["derived"]
            if not isinstance(derived, dict) or set(derived) != {"numerator", "denominator"}:
                raise SourceError(f"{where}: derived needs exactly 'numerator' and 'denominator'")
            for part in derived.values():
                if part not in slugs or part == item["slug"] or "derived" in indicators[slugs.index(part)]:
                    raise SourceError(f"{where}: derived must refer to other, non-derived indicators, not {part!r}")
        seen.add(item["slug"])
    return indicators


def write_indicators_json(path, indicators):
    Path(path).write_text(json.dumps(indicators, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
